import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Colors
    NAVY = RGBColor(27, 58, 107)
    DARK = RGBColor(15, 32, 58)
    WHITE = RGBColor(255, 255, 255)
    LIGHT_BG = RGBColor(246, 248, 252)
    ACCENT = RGBColor(59, 130, 246)
    MUTED = RGBColor(120, 133, 155)
    TEXT = RGBColor(35, 45, 60)
    GREEN = RGBColor(22, 163, 74)
    ORANGE = RGBColor(234, 88, 12)
    CARD_BG = RGBColor(255, 255, 255)
    CARD_BORDER = RGBColor(226, 232, 240)

    def bg(slide, color):
        s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background()

    def textbox(slide, l, t, w, h):
        tb = slide.shapes.add_textbox(l, t, w, h)
        tb.text_frame.word_wrap = True
        return tb.text_frame

    def card(slide, l, t, w, h):
        c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
        c.fill.solid(); c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = CARD_BORDER; c.line.width = Pt(1)

    def add_p(tf, text, size=12, bold=False, color=TEXT, after=6, align=None):
        if len(tf.paragraphs) == 1 and tf.paragraphs[0].text == "":
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = text; p.font.size = Pt(size); p.font.bold = bold
        p.font.color.rgb = color; p.space_after = Pt(after)
        if align: p.alignment = align
        return p

    # =============================================
    # SLIDE 1 — Cover
    # =============================================
    s = prs.slides.add_slide(blank)
    bg(s, DARK)
    # accent bar
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(2.2), Inches(0.08), Inches(2.8))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()

    tf = textbox(s, Inches(1.6), Inches(2.0), Inches(10), Inches(3.5))
    add_p(tf, "REINVENT 4", 14, True, ACCENT, 8)
    add_p(tf, "Web Interface untuk\nGenerative Drug Design", 38, True, WHITE, 16)
    add_p(tf, "Mengubah workflow CLI menjadi aplikasi web interaktif\nyang bisa langsung dipakai peneliti.", 16, False, MUTED, 30)
    add_p(tf, "[Nama]  ·  [NIM]  ·  [Program Studi]  ·  2026", 13, False, RGBColor(160, 175, 200), 0)

    # =============================================
    # SLIDE 2 — Kenapa Dibuat?
    # =============================================
    s = prs.slides.add_slide(blank)
    bg(s, LIGHT_BG)
    tf = textbox(s, Inches(0.9), Inches(0.5), Inches(11), Inches(0.8))
    add_p(tf, "Kenapa Dibuat?", 26, True, NAVY, 4)

    # left: masalah
    card(s, Inches(0.9), Inches(1.5), Inches(5.6), Inches(5.3))
    tf = textbox(s, Inches(1.2), Inches(1.7), Inches(5.0), Inches(5.0))
    add_p(tf, "Masalahnya", 13, True, ORANGE, 10)
    add_p(tf, "REINVENT 4 itu powerful, tapi...", 18, True, NAVY, 14)
    bullets = [
        "Semua operasinya lewat terminal — harus ngetik command panjang.",
        "Konfigurasi pakai file TOML yang formatnya strict banget, salah sedikit langsung error.",
        "Hasilnya cuma file CSV teks mentah, harus buka software lain buat lihat bentuk molekulnya.",
        "Gak ada cara gampang buat pantau proses yang lagi jalan — harus cek log manual.",
        "Buat peneliti kimia yang bukan programmer, ini jadi barrier yang serius.",
    ]
    for b in bullets:
        add_p(tf, "·  " + b, 12, False, TEXT, 8)

    # right: solusi
    card(s, Inches(6.9), Inches(1.5), Inches(5.6), Inches(5.3))
    tf = textbox(s, Inches(7.2), Inches(1.7), Inches(5.0), Inches(5.0))
    add_p(tf, "Solusinya", 13, True, GREEN, 10)
    add_p(tf, "Bikin web app di atasnya.", 18, True, NAVY, 14)
    bullets = [
        "Bungkus seluruh pipeline REINVENT 4 ke dalam satu web app.",
        "Upload file, atur parameter, klik tombol — selesai. Gak perlu terminal.",
        "File TOML di-generate otomatis di backend, user gak perlu tau formatnya.",
        "Hasil langsung ditampilin: tabel, metrik, gambar 2D molekul.",
        "Ada job tracker di sidebar buat pantau proses real-time.",
    ]
    for b in bullets:
        add_p(tf, "·  " + b, 12, False, TEXT, 8)

    # =============================================
    # SLIDE 3 — Tech Stack & Arsitektur
    # =============================================
    s = prs.slides.add_slide(blank)
    bg(s, LIGHT_BG)
    tf = textbox(s, Inches(0.9), Inches(0.5), Inches(11), Inches(0.8))
    add_p(tf, "Tech Stack & Arsitektur", 26, True, NAVY, 4)

    boxes = [
        ("Frontend", "Streamlit (port 8501)", ACCENT, [
            "Tampilan web interaktif langsung dari Python.",
            "Form input: slider, checkbox, file uploader.",
            "Render gambar molekul 2D pakai RDKit.",
            "Bilingual (ID/EN) — bisa switch kapan aja.",
            "Landing page + sidebar navigation.",
        ]),
        ("Backend API", "FastAPI (port 8000)", NAVY, [
            "REST endpoints: /sampling, /transfer_learning, /staged_learning, /preprocessing.",
            "Background tasks — proses berat jalan di belakang layar.",
            "TOML builder otomatis dari parameter user.",
            "Auto-detect nama kolom CSV (case-insensitive).",
            "Serve file download (.prior, .csv, .smi).",
        ]),
        ("Engine", "REINVENT 4 + PyTorch", GREEN, [
            "Dijalankan sebagai subprocess di venv terpisah.",
            "Modul datapipeline: filter & normalisasi SMILES.",
            "Neural network generatif (RNN) untuk sampling molekul.",
            "RL optimizer (DAP strategy) untuk optimasi reward.",
            "CPU mode — bisa jalan di laptop biasa.",
        ]),
        ("Storage", "File-based Jobs", MUTED, [
            "Setiap job punya folder sendiri: /jobs/{uuid}/",
            "status.json — tracking status (QUEUED, RUNNING, COMPLETED, FAILED).",
            "run.log — full log proses untuk debugging.",
            "Dataset, model, dan hasil tersimpan per-job.",
            "Data persist meskipun browser di-refresh.",
        ]),
    ]

    for i, (title, subtitle, color, items) in enumerate(boxes):
        left = Inches(0.9 + i * 3.08)
        card(s, left, Inches(1.5), Inches(2.85), Inches(5.3))
        tf = textbox(s, left + Inches(0.2), Inches(1.7), Inches(2.45), Inches(5.0))
        add_p(tf, title.upper(), 11, True, color, 4)
        add_p(tf, subtitle, 15, True, NAVY, 12)
        for it in items:
            add_p(tf, "·  " + it, 10.5, False, TEXT, 6)

    # =============================================
    # SLIDE 4 — Alur Kerja
    # =============================================
    s = prs.slides.add_slide(blank)
    bg(s, LIGHT_BG)
    tf = textbox(s, Inches(0.9), Inches(0.5), Inches(11), Inches(0.8))
    add_p(tf, "Alur Kerja: Dari Dataset Mentah ke Molekul Baru", 26, True, NAVY, 4)

    steps = [
        ("1", "Pre-processing", "Bersihkan data", ACCENT, [
            "Upload CSV dari database (misal ChEMBL).",
            "Filter: heavy atoms, berat molekul, muatan.",
            "Output: file .smi bersih siap training.",
        ]),
        ("2", "Transfer Learning", "Fokuskan model", NAVY, [
            "Pilih base prior + upload dataset target.",
            "Finetune model biar \"kenal\" kelas molekul kita.",
            "Output: model .prior baru (~23 MB).",
        ]),
        ("3", "Staged Learning (RL)", "Optimasi target", ORANGE, [
            "Set target: QED, rentang berat molekul.",
            "Prior jadi pengawas, Agent (hasil TL) dioptimasi.",
            "Output: checkpoint model teroptimasi.",
        ]),
        ("4", "Sampling", "Panen molekul", GREEN, [
            "Generate ratusan/ribuan molekul baru.",
            "Otomatis hitung Validity %, Uniqueness %.",
            "Visualisasi 2D + Tanimoto similarity.",
        ]),
    ]

    for i, (num, title, desc, color, items) in enumerate(steps):
        left = Inches(0.9 + i * 3.08)
        card(s, left, Inches(1.5), Inches(2.85), Inches(5.3))
        tf = textbox(s, left + Inches(0.2), Inches(1.7), Inches(2.45), Inches(5.0))
        add_p(tf, f"STEP {num}", 11, True, color, 4)
        add_p(tf, title, 17, True, NAVY, 4)
        add_p(tf, desc, 12, False, MUTED, 14)

        # arrow atau divider visual
        line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(0.2), Inches(3.15), Inches(2.45), Inches(0.015))
        line.fill.solid(); line.fill.fore_color.rgb = CARD_BORDER; line.line.fill.background()

        for it in items:
            add_p(tf, "·  " + it, 11.5, False, TEXT, 8)

    # =============================================
    # SLIDE 5 — Fitur: Pre-processing
    # =============================================
    s = prs.slides.add_slide(blank)
    bg(s, LIGHT_BG)
    tf = textbox(s, Inches(0.9), Inches(0.5), Inches(11), Inches(0.8))
    add_p(tf, "Fitur: Pre-processing Dataset", 26, True, NAVY, 4)

    card(s, Inches(0.9), Inches(1.5), Inches(5.6), Inches(5.3))
    tf = textbox(s, Inches(1.2), Inches(1.7), Inches(5.0), Inches(5.0))
    add_p(tf, "Apa yang dilakukan", 14, True, NAVY, 12)
    items = [
        "Pakai modul resmi reinvent.datapipeline di belakang layar.",
        "User upload CSV, atur parameter filter, klik mulai.",
        "Filter berdasarkan jumlah atom berat (min/max) dan berat molekul maksimal.",
        "Opsi netralisasi muatan (uncharge) dan preservasi stereokimia.",
        "Hasil: file .smi yang sudah bersih, bisa langsung dipake buat TL.",
    ]
    for it in items:
        add_p(tf, "·  " + it, 12, False, TEXT, 8)

    card(s, Inches(6.9), Inches(1.5), Inches(5.6), Inches(5.3))
    tf = textbox(s, Inches(7.2), Inches(1.7), Inches(5.0), Inches(5.0))
    add_p(tf, "Yang menarik", 14, True, NAVY, 12)
    items = [
        "Auto-detect kolom SMILES — mau namanya 'SMILES', 'Smiles', atau 'smiles', tetap jalan.",
        "Ini nge-solve bug Polars ColumnNotFoundError yang sering bikin gagal kalau pakai CLI.",
        "Preview 20 SMILES pertama langsung di web setelah proses selesai.",
        "Download hasil .smi dalam satu klik.",
        "[SCREENSHOT HALAMAN PREPROCESSING]",
    ]
    for it in items:
        add_p(tf, "·  " + it, 12, False, TEXT, 8)

    # =============================================
    # SLIDE 6 — Fitur: Transfer Learning
    # =============================================
    s = prs.slides.add_slide(blank)
    bg(s, LIGHT_BG)
    tf = textbox(s, Inches(0.9), Inches(0.5), Inches(11), Inches(0.8))
    add_p(tf, "Fitur: Transfer Learning", 26, True, NAVY, 4)

    card(s, Inches(0.9), Inches(1.5), Inches(5.6), Inches(5.3))
    tf = textbox(s, Inches(1.2), Inches(1.7), Inches(5.0), Inches(5.0))
    add_p(tf, "Cara kerjanya", 14, True, NAVY, 12)
    items = [
        "Pilih base model dari dropdown (reinvent, libinvent, linkinvent) atau upload sendiri.",
        "Upload dataset target dalam format CSV.",
        "Atur epochs dan batch size.",
        "Klik jalankan — model di-finetune di background.",
        "Selesai: download file .prior baru yang sudah \"kenal\" molekul target kita.",
    ]
    for it in items:
        add_p(tf, "·  " + it, 12, False, TEXT, 8)

    card(s, Inches(6.9), Inches(1.5), Inches(5.6), Inches(5.3))
    tf = textbox(s, Inches(7.2), Inches(1.7), Inches(5.0), Inches(5.0))
    add_p(tf, "Fitur tambahan", 14, True, NAVY, 12)
    items = [
        "Validation split otomatis — upload 1 file, sistem split jadi train + valid.",
        "Slider proporsi: 5% sampai 50% (default 20% validasi).",
        "Bisa dimatikan via checkbox — berguna buat LibInvent/LinkInvent yang gak perlu split.",
        "Ada tombol 'Muat Data Contoh' — 30 molekul drug-like siap pakai buat testing.",
        "[SCREENSHOT HALAMAN TRANSFER LEARNING]",
    ]
    for it in items:
        add_p(tf, "·  " + it, 12, False, TEXT, 8)

    # =============================================
    # SLIDE 7 — Fitur: Staged Learning (RL)
    # =============================================
    s = prs.slides.add_slide(blank)
    bg(s, LIGHT_BG)
    tf = textbox(s, Inches(0.9), Inches(0.5), Inches(11), Inches(0.8))
    add_p(tf, "Fitur: Staged Learning (Reinforcement Learning)", 26, True, NAVY, 4)

    card(s, Inches(0.9), Inches(1.5), Inches(5.6), Inches(5.3))
    tf = textbox(s, Inches(1.2), Inches(1.7), Inches(5.0), Inches(5.0))
    add_p(tf, "Setup dual-model", 14, True, NAVY, 12)
    items = [
        "Prior (kolom kiri): model dasar bawaan — jadi \"pagar pembatas\".",
        "Agent (kolom kanan): model hasil TL kita — yang dioptimasi.",
        "Bisa pilih dari dropdown atau upload file .prior sendiri.",
        "Konsep: Prior menjaga agar hasil tetap realistis, Agent dioptimasi menuju target.",
        "Ini sesuai cara kerja asli REINVENT 4 — bukan simplifikasi.",
    ]
    for it in items:
        add_p(tf, "·  " + it, 12, False, TEXT, 8)

    card(s, Inches(6.9), Inches(1.5), Inches(5.6), Inches(5.3))
    tf = textbox(s, Inches(7.2), Inches(1.7), Inches(5.0), Inches(5.0))
    add_p(tf, "Parameter optimasi", 14, True, NAVY, 12)
    items = [
        "Max Steps — berapa iterasi RL yang mau dijalankan.",
        "QED Weight — bobot drug-likeness (slider 0-10).",
        "Target MW Min & Max — rentang berat molekul yang diinginkan.",
        "Scoring pakai Geometric Mean multi-component (bawaan REINVENT).",
        "Output: model checkpoint + log CSV per-step.",
        "[SCREENSHOT HALAMAN RL]",
    ]
    for it in items:
        add_p(tf, "·  " + it, 12, False, TEXT, 8)

    # =============================================
    # SLIDE 8 — Fitur: Sampling & Evaluasi
    # =============================================
    s = prs.slides.add_slide(blank)
    bg(s, LIGHT_BG)
    tf = textbox(s, Inches(0.9), Inches(0.5), Inches(11), Inches(0.8))
    add_p(tf, "Fitur: Sampling & Evaluasi Molekul", 26, True, NAVY, 4)

    card(s, Inches(0.9), Inches(1.5), Inches(5.6), Inches(5.3))
    tf = textbox(s, Inches(1.2), Inches(1.7), Inches(5.0), Inches(5.0))
    add_p(tf, "Generate molekul baru", 14, True, NAVY, 12)
    items = [
        "Pilih model (built-in atau upload custom .prior).",
        "Tentukan jumlah molekul yang mau di-generate.",
        "Klik jalankan — REINVENT generate SMILES baru.",
        "Bisa pakai model bawaan, hasil TL, atau hasil RL.",
        "Ada tombol 'Muat Parameter Contoh' buat quick start.",
    ]
    for it in items:
        add_p(tf, "·  " + it, 12, False, TEXT, 8)

    card(s, Inches(6.9), Inches(1.5), Inches(5.6), Inches(5.3))
    tf = textbox(s, Inches(7.2), Inches(1.7), Inches(5.0), Inches(5.0))
    add_p(tf, "Evaluasi otomatis (RDKit)", 14, True, NAVY, 12)
    items = [
        "Validity % — berapa persen molekul yang valid secara kimia.",
        "Uniqueness % — berapa persen yang unik (deteksi mode collapse).",
        "Tanimoto Similarity — input SMILES referensi, hitung kemiripan.",
        "Grid gambar 2D molekul — 10 pertama langsung ditampilin.",
        "Download hasil CSV lengkap.",
        "[SCREENSHOT HASIL SAMPLING + METRIK]",
    ]
    for it in items:
        add_p(tf, "·  " + it, 12, False, TEXT, 8)

    # =============================================
    # SLIDE 9 — UX: Landing Page, Bilingual, Job Tracker
    # =============================================
    s = prs.slides.add_slide(blank)
    bg(s, LIGHT_BG)
    tf = textbox(s, Inches(0.9), Inches(0.5), Inches(11), Inches(0.8))
    add_p(tf, "UX & Fitur Pendukung", 26, True, NAVY, 4)

    cw = Inches(3.64)
    boxes_data = [
        ("Landing Page", [
            "Welcome screen dengan overview 4 fitur utama.",
            "Pemilih bahasa (Indonesia / English) di halaman awal.",
            "Tombol 'Mulai' / 'Start' untuk masuk.",
            "Desain clean: navy + white, tanpa emoji berlebihan.",
        ]),
        ("Job Tracking", [
            "Live status di sidebar: QUEUED, RUNNING, COMPLETED, FAILED.",
            "Riwayat job: container scrollable di sidebar.",
            "Klik 'Lihat Hasil' — hasil ditampilin di area utama.",
            "Tombol 'Tutup' buat collapse hasil biar tetap rapi.",
        ]),
        ("Developer Experience", [
            "start_all.bat — satu klik nyalain API + UI.",
            "Error message yang jelas, bukan crash mentah.",
            "File terorganisir per-job di folder jobs/{uuid}/.",
            "Arsitektur modular: ui.py, api.py, toml_builder.py, translations.py.",
        ]),
    ]

    for i, (title, items) in enumerate(boxes_data):
        left = Inches(0.9 + i * 3.94)
        card(s, left, Inches(1.5), cw, Inches(5.3))
        tf = textbox(s, left + Inches(0.2), Inches(1.7), cw - Inches(0.4), Inches(5.0))
        add_p(tf, title, 16, True, NAVY, 12)
        for it in items:
            add_p(tf, "·  " + it, 12, False, TEXT, 9)

    # =============================================
    # SLIDE 10 — Demo / Screenshot placeholder
    # =============================================
    s = prs.slides.add_slide(blank)
    bg(s, LIGHT_BG)
    tf = textbox(s, Inches(0.9), Inches(0.5), Inches(11), Inches(0.8))
    add_p(tf, "Demo Aplikasi", 26, True, NAVY, 4)

    # Placeholder boxes for screenshots
    for i, label in enumerate(["Landing Page", "Pre-processing", "Transfer Learning"]):
        left = Inches(0.9 + i * 3.94)
        card(s, left, Inches(1.5), Inches(3.64), Inches(5.3))
        tf = textbox(s, left + Inches(0.3), Inches(3.5), Inches(3.04), Inches(1.0))
        add_p(tf, f"[Screenshot: {label}]", 14, False, MUTED, 4, PP_ALIGN.CENTER)

    # =============================================
    # SLIDE 11 — Demo / Screenshot placeholder 2
    # =============================================
    s = prs.slides.add_slide(blank)
    bg(s, LIGHT_BG)
    tf = textbox(s, Inches(0.9), Inches(0.5), Inches(11), Inches(0.8))
    add_p(tf, "Demo Aplikasi (lanjutan)", 26, True, NAVY, 4)

    for i, label in enumerate(["Staged Learning (RL)", "Sampling + Hasil", "Job History"]):
        left = Inches(0.9 + i * 3.94)
        card(s, left, Inches(1.5), Inches(3.64), Inches(5.3))
        tf = textbox(s, left + Inches(0.3), Inches(3.5), Inches(3.04), Inches(1.0))
        add_p(tf, f"[Screenshot: {label}]", 14, False, MUTED, 4, PP_ALIGN.CENTER)

    # =============================================
    # SLIDE 12 — Penutup
    # =============================================
    s = prs.slides.add_slide(blank)
    bg(s, DARK)

    tf = textbox(s, Inches(1.5), Inches(2.2), Inches(10.33), Inches(3.0))
    add_p(tf, "Terima Kasih", 42, True, WHITE, 12, PP_ALIGN.CENTER)
    add_p(tf, "Ada pertanyaan? Atau mau langsung demo?", 18, False, MUTED, 24, PP_ALIGN.CENTER)
    add_p(tf, "localhost:8501  ·  localhost:8000/docs", 13, False, RGBColor(100, 120, 150), 0, PP_ALIGN.CENTER)

    prs.save(output_path)
    print(f"Done: {output_path}")

create_deck("d:/S2/REINVENT/web_server/Presentasi_REINVENT4.pptx")
